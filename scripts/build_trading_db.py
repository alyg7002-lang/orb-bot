"""Build data/trading_knowledge.db (SQLite) + reports/trading_knowledge.xlsx.

Tables: glossary, strategies (orb/knowledge.py), indicators and candlestick_patterns
(every function in TA-Lib, the standard open-source technical analysis library).
    pip install TA-Lib openpyxl
    python scripts/build_trading_db.py
"""
import sqlite3, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import talib, talib.abstract as ta
from orb.knowledge import GLOSSARY, STRATEGIES

ROOT = Path(__file__).resolve().parents[1]
db_path = ROOT / "data" / "trading_knowledge.db"
xl_path = ROOT / "reports" / "trading_knowledge.xlsx"

ind, pat = [], []
for group, names in talib.get_function_groups().items():
    for n in names:
        i = ta.Function(n).info
        inputs = ", ".join(sorted({x for v in i["input_names"].values() for x in (v if isinstance(v, list) else [v])}))
        params = ", ".join(f"{k}={v}" for k, v in i["parameters"].items())
        row = (n, i["display_name"], group, inputs, params, ", ".join(i["output_names"]))
        (pat if group == "Pattern Recognition" else ind).append(row)

tables = {
    "glossary": (["term", "category", "definition"], GLOSSARY),
    "strategies": (["name", "style", "timeframe", "source", "rules", "testable_by_bot"], STRATEGIES),
    "indicators": (["talib_name", "name", "category", "inputs", "default_params", "outputs"], ind),
    "candlestick_patterns": (["talib_name", "name", "category", "inputs", "default_params", "outputs"], pat),
}

db_path.unlink(missing_ok=True)
con = sqlite3.connect(db_path)
for t, (cols, rows) in tables.items():
    con.execute(f"CREATE TABLE {t} ({', '.join(c + ' TEXT' for c in cols)})")
    con.executemany(f"INSERT INTO {t} VALUES ({','.join('?' * len(cols))})", rows)
con.commit(); con.close()

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
wb = Workbook(); wb.remove(wb.active)
widths = {"definition": 90, "rules": 80, "source": 50, "name": 38, "term": 34}
for t, (cols, rows) in tables.items():
    ws = wb.create_sheet(t.replace("_", " ").title())
    ws.append([c.replace("_", " ").title() for c in cols])
    for r in rows:
        ws.append(list(r))
    for c, col in zip(ws[1], cols):
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F")
        ws.column_dimensions[c.column_letter].width = widths.get(col, 20)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
xl_path.parent.mkdir(exist_ok=True)
wb.save(xl_path)
print({t: len(r) for t, (_, r) in tables.items()}, db_path, xl_path)
